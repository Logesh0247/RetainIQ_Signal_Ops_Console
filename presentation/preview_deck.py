"""
Render the generated deck to PNG previews.

There is no PowerPoint/LibreOffice in this environment, so this draws the
shapes and text of each slide directly with Pillow. It is an approximation of
PowerPoint rendering (same geometry, same colours, substituted fonts), but it
is close enough to judge layout, balance and density visually.

    python presentation/preview_deck.py [slide_numbers...]
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
DECK = HERE / "RetainIQ_Project_Review.pptx"
OUT = HERE / "preview"
OUT.mkdir(exist_ok=True)

SCALE = 1600 / 13.333  # px per inch
W, H = int(13.333 * SCALE), int(7.5 * SCALE)

_TTF = Path(__import__("matplotlib").__file__).parent / "mpl-data" / "fonts" / "ttf"
FONTS = {
    ("body", False): str(_TTF / "DejaVuSans.ttf"),
    ("body", True): str(_TTF / "DejaVuSans-Bold.ttf"),
    ("mono", False): str(_TTF / "DejaVuSansMono.ttf"),
    ("mono", True): str(_TTF / "DejaVuSansMono-Bold.ttf"),
}
_cache: dict = {}


def font_for(mono: bool, bold: bool, pt: float):
    key = ("mono" if mono else "body", bold, round(pt, 1))
    if key not in _cache:
        px = max(6, int(round(pt / 72 * SCALE)))   # points -> inches -> pixels
        _cache[key] = ImageFont.truetype(FONTS[(key[0], bold)], px)
    return _cache[key]


def measure(text, f):
    box = f.getbbox(text)
    return box[2] - box[0], box[3] - box[1]


def wrap(text, f, max_px):
    words, lines, cur = text.split(), [], ""
    for word in words:
        trial = f"{cur} {word}".strip()
        if measure(trial, f)[0] <= max_px or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def rgb(color):
    try:
        return (color.rgb[0], color.rgb[1], color.rgb[2])
    except Exception:
        return (255, 255, 255)


def shape_fill(shape):
    try:
        if shape.fill.type is not None and shape.fill.type == 1:
            return rgb(shape.fill.fore_color)
    except Exception:
        pass
    return None


def shape_line(shape):
    try:
        if shape.line.fill.type == 1:
            return rgb(shape.line.color), max(1, int(round((shape.line.width or Emu(9525)) / 914400 * SCALE)))
    except Exception:
        pass
    return None, 0


def draw_rounded(d, box, radius, fill, outline, width):
    if fill is None and outline is None:
        return
    d.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def render(slide, index):
    img = Image.new("RGB", (W, H), (0x0A, 0x0F, 0x1C))
    d = ImageDraw.Draw(img)

    for shape in slide.shapes:
        x, y = shape.left / 914400 * SCALE, shape.top / 914400 * SCALE
        w, h = shape.width / 914400 * SCALE, shape.height / 914400 * SCALE
        st = shape.shape_type

        # picture
        if st == 13:
            try:
                blob = shape.image.blob
                import io
                pic = Image.open(io.BytesIO(blob)).convert("RGB")
                pic = pic.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
                img.paste(pic, (int(x), int(y)))
            except Exception:
                pass
            continue

        is_auto = False
        try:
            shape.auto_shape_type
            is_auto = True
        except Exception:
            is_auto = False

        if is_auto:
            fill = shape_fill(shape)
            outline, lw = shape_line(shape)
            radius = 10
            try:
                adj = shape.adjustments[0]
                radius = max(0, int(adj * min(w, h)))
            except Exception:
                radius = 0
            if radius == 0:
                d.rectangle((x, y, x + w, y + h), fill=fill, outline=outline, width=lw)
            else:
                draw_rounded(d, (x, y, x + w, y + h), radius, fill, outline, lw)

        if shape.has_text_frame and shape.text_frame.text.strip():
            tf = shape.text_frame
            cx = x
            cy = y
            for para in tf.paragraphs:
                runs = para.runs
                if not runs:
                    continue
                pt = max((r.font.size.pt for r in runs if r.font.size), default=12)
                bold = any(r.font.bold for r in runs)
                mono = any((r.font.name or "").lower().startswith("consolas") for r in runs)
                color = rgb(runs[0].font.color) if runs[0].font.color and runs[0].font.color.type is not None else (255, 255, 255)
                text = "".join(r.text for r in runs)
                f = font_for(mono, bool(bold), pt)

                max_px = max(10, w - 2)
                lines = wrap(text, f, max_px)

                spacing = para.line_spacing or 1.12
                line_h = pt / 72 * SCALE * 1.22 * (spacing if isinstance(spacing, float) else 1.12)

                align = str(para.alignment)
                for ln in lines:
                    lw_, lh_ = measure(ln, f)
                    lx = cx
                    if "RIGHT" in align:
                        lx = cx + w - lw_
                    elif "CENTER" in align:
                        lx = cx + (w - lw_) / 2
                    d.text((lx, cy), ln, font=f, fill=color)
                    cy += line_h
                sa = para.space_after
                if sa is not None:
                    cy += sa.pt / 72 * SCALE
                sb = para.space_before
                if sb is not None:
                    cy += sb.pt / 72 * SCALE

    path = OUT / f"slide_{index:02d}.png"
    img.save(path)
    return path


def main():
    wanted = [int(a) for a in sys.argv[1:]] if len(sys.argv) > 1 else None
    prs = Presentation(str(DECK))
    made = []
    for i, slide in enumerate(prs.slides, 1):
        if wanted and i not in wanted:
            continue
        made.append(render(slide, i))
    print(f"rendered {len(made)} slide preview(s) to {OUT.relative_to(HERE.parent)}/")


if __name__ == "__main__":
    main()
